"""Golden-master tests for xlsx2csv.format_cell_value() and convert_xlsx_to_csv()."""

import json
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from scripts.work_with_csv.xlsx2csv import convert_xlsx_to_csv, format_cell_value

FIXTURES_PATH = Path(__file__).parent / "test_xlsx2csv_fixtures.json"
FIXTURES = json.loads(FIXTURES_PATH.read_text(encoding="utf-8"))


def _build_input_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "analysis"

    ws.append(["plain", "bold_cell", 5.0, 5.5, 7, None])
    ws["B1"].font = ws["B1"].font.copy(b=True)

    rt = CellRichText(
        TextBlock(InlineFont(b=True), "bold part"),
        " plain part",
    )
    ws["A2"] = rt

    wb.save(path)


def test_format_cell_value_matches_fixture(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "input.xlsx"
    _build_input_workbook(xlsx_path)

    wb = openpyxl.load_workbook(xlsx_path, rich_text=True, data_only=True)
    sheet = wb["analysis"]

    expected = {
        c["coordinate"]: c["result"] for c in FIXTURES["format_cell_value"]["cells"]
    }
    for row in sheet.iter_rows():
        for cell in row:
            assert isinstance(cell, Cell)
            assert format_cell_value(cell) == expected[cell.coordinate]


def test_convert_xlsx_to_csv_matches_fixture(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "input.xlsx"
    csv_path = tmp_path / "output.csv"
    _build_input_workbook(xlsx_path)

    convert_xlsx_to_csv(xlsx_path, csv_path, "analysis")

    assert (
        csv_path.read_text(encoding="utf-8")
        == FIXTURES["convert_xlsx_to_csv"]["output_content"]
    )


def test_convert_xlsx_to_csv_missing_input_file_does_not_raise(tmp_path: Path) -> None:
    csv_path = tmp_path / "output.csv"
    convert_xlsx_to_csv(tmp_path / "does_not_exist.xlsx", csv_path, "analysis")
    assert not csv_path.exists()


def test_convert_xlsx_to_csv_missing_sheet_does_not_raise(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "input.xlsx"
    csv_path = tmp_path / "output.csv"
    _build_input_workbook(xlsx_path)

    convert_xlsx_to_csv(xlsx_path, csv_path, "no_such_sheet")
    assert not csv_path.exists()
